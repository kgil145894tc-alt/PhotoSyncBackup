from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(r'C:\Users\kirlg\PhotoSync')
OUT=ROOT/'tmp/defense-guide'
OUT.mkdir(parents=True,exist_ok=True)
d=Document()
s=d.sections[0]; s.top_margin=Inches(.7); s.bottom_margin=Inches(.65); s.left_margin=s.right_margin=Inches(.8)
s.page_width=Inches(8.5); s.page_height=Inches(11)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
 st=d.styles[name]; st.font.name='Calibri'; st.font.color.rgb=RGBColor(0,0,0)
d.styles['Normal'].font.size=Pt(11)
d.styles['Normal'].paragraph_format.space_after=Pt(7)
d.styles['Normal'].paragraph_format.line_spacing=1.08
d.styles['Title'].font.size=Pt(28)
d.styles['Heading 1'].font.size=Pt(20)
d.styles['Heading 2'].font.size=Pt(13)
footer=s.footer.paragraphs[0]; footer.alignment=2
footer.add_run('PhotoSync Defense Guide  |  ')
f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');footer._p.append(f)
def p(t): d.add_paragraph(t)
def h(t):d.add_heading(t,2)
def page(t):
 d.add_page_break();d.add_heading(t,1)
def bullets(items):
 for x in items:d.add_paragraph(x,'List Bullet')
def table(headers,rows,widths=None):
 t=d.add_table(rows=1,cols=len(headers));t.autofit=False
 if widths:
  for c,w in zip(t.columns,widths):c.width=Inches(w)
 for c,v in zip(t.rows[0].cells,headers):c.text=v
 for row in rows:
  for c,v in zip(t.add_row().cells,row):c.text=str(v)
 for i,row in enumerate(t.rows):
  for cell in row.cells:
   pr=cell._tc.get_or_add_tcPr()
   borders=OxmlElement('w:tcBorders')
   for side in ['top','left','bottom','right']:
    e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
   pr.append(borders)
   sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'243E59' if i==0 else ('F2F5F8' if i%2==0 else 'FFFFFF'));pr.append(sh)
   for para in cell.paragraphs:
    para.paragraph_format.space_after=Pt(5);para.paragraph_format.space_before=Pt(5)
    for run in para.runs:
     run.font.size=Pt(9.5)
     if i==0:run.font.bold=True;run.font.color.rgb=RGBColor(255,255,255)
  if i==0:
   rep=OxmlElement('w:tblHeader');row._tr.get_or_add_trPr().append(rep)
 p('')

d.add_heading('PhotoSync Project Defense Guide',0)
d.add_paragraph('Implementation details and presentation preparation',style='Subtitle')
p('Prepared October 9 2026  |  Defense October 10 2026')
p('PhotoSync is a mobile photography studio booking system. Clients browse services and packages, request an available session, and track their bookings. Photographers use the admin workspace to review requests, manage availability and the catalog, and maintain studio information.')
p('The central technical design is an Expo and React Native application connected to Supabase authentication, PostgreSQL, storage, database functions, and a push notification Edge Function. The server controls access and protects confirmed appointments; the app supplies the booking workflow and responsive screens.')
h('How to study this guide')
p('Start with the overview and backend sections, then rehearse the client and photographer workflows. Use the limitations section to avoid claiming prototype or planned features as completed work. The questions and demonstration checklist are designed for presentation practice.')
table(['Section','What you should be able to explain'],[
('1 Project overview','Purpose, users, scope and technology choices'),('2 Backend architecture','Supabase services and request flow'),('3 Database design','Entities, relationships and constraints'),('4 Backend operations','APIs, database functions and notifications'),('5 Authentication and security','Roles, sessions, RLS and storage access'),('6 Booking rules','Validation, conflicts and lifecycle'),('7 Mobile application','Screens, routing and code organization'),('8 Performance and reliability','Caches, paging, concurrency and errors'),('9 Setup and deployment','Environment, SQL setup and builds'),('10 Validation and limitations','Current checks and remaining verification'),('11 Defense questions','Suggested answers grounded in the code'),('12 Demonstration plan','A practical rehearsal sequence'),('Appendices','Source map and SQL routine inventory')],[1.85,5.05])

page('1 Project overview')
h('Problem and proposed solution')
p('A photography studio needs a consistent way to collect session requests, organize package choices, communicate booking decisions and control available times. PhotoSync combines these tasks in one client and photographer application. This problem statement describes the implemented workflow; it is not a claim that a user study or measured business improvement has been completed.')
h('Users and responsibilities')
bullets(['Client: create an account, sign in, browse published services and packages, select a date and time, provide contact and session details, submit a request, view booking status, cancel a pending request or request rescheduling, and manage a profile.','Photographer or admin: review requests, confirm or reject pending bookings, complete confirmed sessions, manage calendar availability, create or edit services and packages, manage catalog images and publication, and edit studio details.','Supabase backend: authenticate users, persist shared records, enforce row access and database constraints, execute database functions, and support reminder and push workflows.'])
h('Technology stack from package configuration')
table(['Technology','Declared version or implementation','Purpose'],[
('Expo','~57.0.25','Mobile framework and build tooling'),('React Native','0.86.3','Native application UI'),('React','19.2.3','Component model'),('Expo Router','~57.0.23','File based navigation'),('TypeScript','~6.0.3','Static checks and shared types'),('Supabase JavaScript','^2.117.1','Auth, database, RPC and storage client'),('PostgreSQL','Hosted through Supabase','Relational data and business constraints'),('Deno Edge Function','send-push-notification','Server side push dispatch'),('AsyncStorage','2.2.0','Configured auth persistence adapter')],[1.5,2.1,3.3])
h('What the project name means in this implementation')
p('Explain PhotoSync as coordination of photography bookings and studio information. The main project does not implement automatic camera roll synchronization, cloud photo backups, an image editing pipeline or delivery of completed photo galleries. Its image uploads support profiles and the service catalog.')
h('Separate prototype')
p('The sibling PhotoSyncMobile project is a separate local booking prototype. Its README describes SQLite or localStorage persistence, demo account switching, client verification and payment status tracking. These do not establish the same features in the main Supabase backed PhotoSync application.')

page('2 Backend architecture')
h('Backend as a service')
p('PhotoSync does not contain a separate Express or custom Node REST server. Application services call Supabase directly through a shared client in src/lib/supabase.ts. Supabase provides Auth, a PostgreSQL database exposed through its data API, object storage and database RPCs. A Deno Edge Function handles push delivery.')
h('Request path')
bullets(['A screen calls a hook or service in src/hooks or src/services.','The service obtains the current authenticated user when required, validates its inputs and invokes a table query or RPC through the Supabase client.','Supabase evaluates the caller identity and PostgreSQL row level security policies. Constraints and database functions apply additional rules.','The service maps database rows into application types and returns a success, error or data result. Stores update cached state and publish change events to other screens.'])
h('Main backend components')
table(['Component','Responsibility'],[('Supabase Auth','Password login, signup, recovery, Google token exchange and sessions'),('PostgreSQL tables','Profiles, catalog, bookings, availability, notifications and history'),('RLS policies','Client ownership and admin permissions'),('Database functions','Availability, booking pagination, expiration, summaries and reminders'),('Supabase Storage','Profile avatars and service or package images'),('Database schedule','Hourly generation of next day booking reminders'),('Edge Function','Read a notification and send it to active device sessions'),('Expo Push Service and FCM configuration','Mobile push delivery infrastructure')],[2.1,4.8])
h('Typical booking request')
p('Submission validates the draft, checks authentication, reloads the selected package and service rules, verifies availability and writes contact information. It inserts a pending booking, then attempts status history and admin notifications. Other booking views are invalidated through local change events. Submission is a request for approval, not an immediate confirmed reservation.')
h('Architecture tradeoff')
p('Direct Supabase access reduces custom server code and gives the application shared data and authentication quickly. It makes correct database policies and migrations essential. A hidden admin button is not a permission boundary; the backend must reject unauthorized calls regardless of the screen used.')

page('3 Database design')
p('UUID identifiers connect business records. PostgreSQL stores structured data; storage buckets hold image files and tables retain their URLs. The base schema and feature SQL files together describe the intended database.')
table(['Entity','Important fields','Relationship or purpose'],[
('profiles','id, role, email, full_name, phone, username, avatar_url','id references auth.users; client or admin role'),('services','id, slug, name, duration_minutes, buffer_minutes, minimum_notice_days, price, is_active, archived_at, image_url','Parent category for packages and scheduling rules'),('packages','id, service_id, name, price, badge, inclusions, is_active, archived_at, image_url','Many packages belong to one service'),('bookings','id, client_id, service_id, package_id, booking_date, start_time, end_time, status','Each booking links a client and selected offering'),('time_slots','id, slot_date, start_time, end_time, status, reason','Admin available or unavailable windows'),('notifications','id, user_id, booking_id, title, message, is_read, created_at','Inbox items for a recipient and optional booking'),('booking_status_history','id, booking_id, status, changed_by, reason, metadata','Timeline of status events'),('audit_logs','id, actor_id, action, entity_type, entity_id, metadata','Record of business actions'),('studio_settings','id, studio_name, studio_address, contact_phone, contact_email, business_hours','One studio configuration row'),('push_tokens','Device token and session association','Device registrations used for delivery')],[1.45,2.8,2.65])
h('Booking detail fields')
p('Bookings also store contact_name, contact_email, contact_phone, notes, people_count, shoot_location, session_theme and special_requests. A rejection_reason field is added through the rejection migration. Contact fields preserve request details separately from the user profile.')
h('Relationships to explain verbally')
p('One authenticated user has a profile. A client profile can have many bookings. A service can have many packages; bookings reference the service and package. A booking can have many status history entries and notifications. A recipient profile can have many notifications and device registrations. Calendar slots describe studio time rather than belonging to a particular client.')
h('Important constraints')
p('The schema restricts roles and booking statuses to known values. Time slots have a uniqueness rule for date, start and end. Usernames have a case insensitive unique index. The studio settings row is a singleton. A GiST exclusion constraint prevents overlapping confirmed booking intervals, using half open ranges so adjacent appointment endpoints can meet. This constraint alone does not enforce preparation buffers.')

page('4 Backend operations')
h('API organization')
p('The app primarily uses Supabase Auth methods, generated table endpoints and PostgreSQL RPCs. It does not define a custom /login or /bookings REST controller. Examples of generated resource paths are /rest/v1/bookings and /rest/v1/rpc/get_booking_page; the project URL supplies the host.')
table(['Operation','Code or database entry point','Behavior'],[
('Authentication','auth.signInWithPassword, signUp, resetPasswordForEmail','Login, registration and password recovery'),('Catalog browsing','services and packages queries','Published records and fresh package validation'),('Create request','bookings insert','New request starts pending'),('Booking lists','get_booking_page','30 record pages with totals and cursor'),('Availability','get_active_booking_slots','Busy intervals without exposing full client profiles'),('Calendar summary','get_calendar_day_summaries','Month aggregate instead of fetching every booking'),('Expire old requests','expire_past_pending_bookings','Server clock and transactional status history'),('Archive catalog','archive_catalog_item','Archive record while preserving booking references'),('Admin notifications','create_admin_booking_notification','Notify admins without reading their profiles'),('Read inbox','notifications query and read RPCs','Recipient filtering and read state updates'),('Reminders','Scheduled and authenticated routines','Tomorrow confirmed bookings and duplicate prevention'),('Push delivery','send-push-notification Edge Function','Notification ID to active recipient devices')],[1.3,2.8,2.8])
h('Push delivery sequence')
p('The device registers an Expo token for its signed in session. A notification row is created. A configured webhook calls the Edge Function with notification_id or record.id. The function loads the notification with a server credential, retrieves active tokens through get_active_push_tokens_for_user, and submits messages to Expo Push Service. The payload includes booking, user, notification and session IDs plus a route.')
p('The function accepts POST and handles OPTIONS. It rejects malformed IDs, missing notifications and failed upstream calls with HTTP errors. PUSH_WEBHOOK_SECRET is checked when configured. The server credential belongs only in the function environment. The returned sent value counts messages submitted, not proof that every phone displayed a push.')
h('Reminders')
p('The documented hourly job selects tomorrow confirmed sessions in Asia/Manila time. A partial unique index and conflict handling prevent repeated reminder inserts. The authenticated app reminder routine remains recipient scoped. Reminder generation and device delivery are separate stages.')

page('5 Authentication and security')
h('Account flow')
p('Clients register a full name, email, username and password. The app normalizes email and username, checks required fields and email shape, and requires at least six password characters. A username lookup RPC supports username login. Supabase Auth owns password authentication; passwords are not stored in the profiles table.')
p('A database signup trigger creates the client profile and assigns the client role. Role lookup routes admins to /photographer and clients to /home. Admin assignment is a trusted database operation, not a choice offered at public signup. The root auth routing and route guards coordinate session restoration and redirects.')
h('Google and password recovery')
p('Android uses a native Google chooser and exchanges its ID token with Supabase. A fresh nonce ties the attempt to its verification, with SHA-256 sent to Google and the original nonce to Supabase. The generic platform implementation returns that Google sign-in is available in Android. Recovery supports a code exchange or access and refresh tokens from a reset link, then calls updateUser for the new password.')
h('Access controls')
bullets(['Clients read their own profiles and bookings; admins can read broader operational records under admin policies.','Catalog write and availability management permissions are restricted to admins. Published catalog records have read policies.','Notification inbox access is recipient scoped. Dedicated RPCs handle operations that should not require a client to read admin profiles.','Profile updates are intended to preserve role, and profile creation requires the caller ID and client role.','Database functions must be examined individually: SECURITY INVOKER retains caller privileges; SECURITY DEFINER executes with its owner privileges and needs explicit authorization and grants.'])
h('Image storage and privacy')
p('The profile-images and service-images buckets are configured as public. Profile writes are restricted by a user ID folder prefix; catalog image writes are admin only. Public reading means these assets should not be described as private encrypted galleries. Database ownership controls and public object visibility are separate concepts.')
h('Sessions and secrets')
p('The client is configured for session persistence and token refresh, using AsyncStorage when the runtime provides window and a no-op storage adapter otherwise. This is not evidence of hardware backed secure storage. Logout attempts push registration cleanup and clears account scoped caches; late results are guarded so they cannot repopulate another account. Public Expo environment variables contain the project URL and anon key. Service role keys, webhook secrets, Google client secrets and private FCM credentials must stay server side.')

page('6 Booking rules and lifecycle')
table(['Current status','Allowed action in application','Result'],[
('New draft','Client submits valid request','pending'),('pending','Admin confirms after validation','confirmed'),('pending','Admin rejects with reason handling','rejected'),('pending','Client cancels own request','cancelled'),('pending','Scheduled start has elapsed and expiration RPC runs','expired'),('pending or confirmed','Client requests valid reschedule','pending for new review'),('confirmed','Admin marks finished session','completed')],[1.4,3.6,1.9])
h('Submission validation')
bullets(['Require a selected package, schedule and contact information; validate email and phone.','Reject sample package IDs rather than silently writing a fake booking.','Reload the actual package and parent service. Reject unavailable offerings, and require review again when the price changed.','Check the slot fits service duration, minimum notice, availability windows and preparation buffer. The submission helper applies at least one day notice.','Check busy booking intervals and unavailable calendar windows. When custom open windows do not exist, load studio working hours.','Require successful server verification before saving, then insert the request as pending.'])
h('Conflict example')
p('Suppose a service lasts 60 minutes and requires a 30 minute buffer. A requested appointment must fit the opening window with its trailing preparation allowance and avoid conflicts when its buffered interval is compared with occupied or unavailable times. The exact code determines how boundaries are applied; do not equate the simple confirmed interval exclusion constraint with every application buffer rule.')
h('Concurrency protection')
p('Admin status updates include the expected previous status in the database filter. Confirm and reject require pending; completion requires confirmed. If another action already changed the row, the app reports that it must be refreshed. If two confirmations race, the database confirmed overlap constraint provides a final protection against overlapping confirmed appointments.')
h('Expiration and transaction boundaries')
p('Expiration uses the server clock in Asia/Manila and updates due pending rows plus history in one transaction. Repeated calls do not create duplicate transitions. Other workflows often write the booking first and write history, audit or notifications afterward through separate requests. Do not claim that the entire submission or every admin action is one atomic database transaction.')
h('Calendar time protection')
p('Past calendar dates remain visible but are read only. Today slots must start after the current studio time, except the supported full day unavailable marker. App checks are repeated before writes, and a database trigger supplies protection against bypassing the UI. Some booking date helpers still use device local dates, which is a timezone consistency point to test.')

page('7 Mobile application')
h('Code organization')
table(['Location','Responsibility'],[('src/app','Route screens and navigator layouts'),('src/components','Shared UI, forms, alerts, navigation and inbox views'),('src/hooks','Screen lifecycle and data subscriptions'),('src/services','Supabase operations, rules, caches, stores and events'),('src/types','Auth, booking, catalog and calendar contracts'),('src/styles','Screen styling, themes and responsive rules'),('src/navigation','Tab and scroll navigation and motion helpers'),('docs','Database scripts and implementation notes'),('supabase/functions','Server push dispatcher'),('scripts','Regression tests and SQL verification fixtures')],[1.8,5.1])
h('Client route flow')
p('Public entry screens include index, login, create-account and reset-password. The client group contains home, services, profile, notifications, book and about. Services have dynamic slug pages and a portrait route. The booking path contains selected, schedule, information, review and success screens. Profile editing is a separate route.')
p('A typical client selects a package, chooses date and time, enters contact and session details, reviews the draft, submits, and sees a pending confirmation. The booking draft service keeps selections in module memory. A full application restart is not demonstrated to restore an unfinished draft.')
h('Photographer route flow')
p('The photographer group contains the dashboard, requests list, request detail by ID, calendar, calendar slots, services, notifications, profile and Clients screen. The dashboard and request list use paged booking data. Detail views load a specific booking and coordinate decision actions. Calendar screens manage studio availability. Service forms manage categories, packages, publication and images.')
h('Mobile interaction choices')
p('Shared keyboard aware form scrolling helps fields remain visible. Safe area and navigation height helpers avoid placing content under system or app navigation. The project has native and web navigation implementations, responsive styles, reduced motion handling and screen transition components. Catalog screens use FlatList virtualization. These are implementation choices; performance on every device still requires physical testing.')
h('Image handling')
p('Image picker integration supports selecting avatars and catalog images. Storage services upload bytes and save a public URL into the related record. Bundled assets supply branding and display defaults. There is no implemented workflow here for uploading and delivering a client photoshoot album.')

page('8 Performance and reliability')
h('Why caching is present')
p('Returning to tabs should not reload the same data unnecessarily or replace useful content with a blank loading screen. Caches retain successful snapshots while refreshing and distinguish an empty response from a failed request. Stores and events notify screens when changes invalidate their data.')
table(['Mechanism','Implementation detail','Benefit'],[
('Booking pagination','30 rows and a created time plus ID cursor','Limits each list response with stable ordering'),('Catalog snapshot paging','Stable 200 row reads until complete','Avoids truncation at the data API row cap'),('Catalog screen virtualization','Eight initial or batch items and seven viewport window','Limits rendered cards'),('Home highlights','12 row reads until three valid highlights','Narrow landing page data'),('Catalog and settings reuse','Common 60 second freshness windows','Reduces repeated browsing reads'),('Expiration and reminder coordination','30 second successful read freshness','Joins repeated maintenance calls'),('Calendar month summaries','At most one aggregate row per date','Avoids downloading all monthly bookings'),('Admin detail cache','At most 24 idle detail entries','Bounds retained idle records')],[1.65,2.75,2.5])
h('Fresh validation differs from cached presentation')
p('Booking submission reloads package data. Confirmation, cancellation, rescheduling and related mutations force expiration checks. Working hours checks bypass browsing freshness when used to authorize a write. Failed verification stops the write. A cached screen improves browsing but is not proof that a booking remains available.')
h('Concurrent requests and stale responses')
p('Consumers can share an unfinished read. Account generations and request revisions reject late responses after logout, account changes or invalidation. A confirmed save can update the shared cache without another read. Older reads cannot overwrite a newer accepted update. Failures are not cached as successful results.')
h('Boundaries')
p('Catalog snapshots still read the complete catalog in pages; virtualization reduces UI work, not the total catalog snapshot size. Caches are not an offline synchronization engine. Local change events coordinate screens within this app process and are not proof of a realtime cross-device subscription. Pull refresh, focus, resume and server reads remain important for fresh information.')

page('9 Setup and deployment')
h('Local prerequisites')
p('Use a supported Node and npm setup with the repository package-lock.json. This checkout does not contain a bun.lock. Install dependencies with npm install and use the repository scripts. The postinstall script patches Expo Router linking behavior and should be understood before troubleshooting navigation startup.')
h('Environment')
p('Copy .env.example to .env and supply EXPO_PUBLIC_SUPABASE_URL and EXPO_PUBLIC_SUPABASE_ANON_KEY. Restart the development server after changing environment variables. Configure Google provider values and signing certificates in their respective services. Do not copy private credentials into presentation material.')
h('Database setup order')
p('Use docs/supabase-setup.md as the detailed checklist. Start with the base schema and catalog seed, then apply catalog archives, contact fields, calendar support and date guards, summaries, buffer and minimum notice, booking actions and status migrations, expiration, studio settings, multiple active bookings, image storage, notifications and RPCs, reminders, read state, usernames and push support. Review hardening, conflict protection and pagination scripts as well.')
p('Do not assume every SQL file is standalone or safe to run repeatedly. In particular, the base schema history backfill references rejection_reason while that column is introduced in a separate migration. A new database setup should be checked for dependencies rather than applying files blindly. Existing repository notes describe past deployments, not a fresh live audit performed for this guide.')
h('Useful project commands')
bullets(['npm install','npx expo start --dev-client','npx expo lint','npx tsc --noEmit','node --test scripts/*.test.cjs','npx expo-doctor','npx expo install <package> for SDK compatible package installation'])
h('Build configuration')
p('app.json sets portrait orientation, the photosync scheme, the Android package com.kirl123.PhotoSync and static web output. Plugins include Router, notifications, image picker, splash and native Google sign-in. Android Firebase configuration points to google-services.json. The checkout contains an android directory; do not describe this repository as having no native project.')
p('eas.json declares development, preview and production profiles. Development uses a development client and internal distribution; preview is internal; production auto increments the build version. Example build: npx eas-cli@latest build --profile development --platform android. Production submission is configured as a profile, but configuration is not proof of a published store release. Native Google and push verification require an appropriate native build.')

page('10 Validation and limitations')
h('Checks run for this guide')
p('The application regression command completed with 473 tests passed, zero failed, zero skipped and zero cancelled. TypeScript completed successfully with no reported errors. Expo lint also completed successfully with no reported errors or warnings. These checks validate code behavior in their test environment; they do not certify deployed backend settings or physical device push delivery.')
h('What the tests cover')
p('The scripts include auth routing and signup, Google token flow, logout, session races, booking pagination and actions, package validation, expiration, calendar rules, service catalog, settings, notification counts and inbox state, push dispatch, cache lifetimes and UI wiring. Several tests use mocks or source level assertions. SQL fixture scripts are separate from the JavaScript test command.')
h('Earlier verification records')
p('Repository notes dated October 7 describe deployment and rollback fixture verification for calendar summaries, reminder generation, expiration, pagination and date guards. The notes report an active hourly reminder job. Those historical records support the project history; the live database and scheduled deliveries were not rechecked during this documentation task.')
h('Known scope limitations')
bullets(['The photographer Clients screen renders placeholder sample entries; it is not a completed live customer directory.','No payment gateway, payment status database workflow, invoice generation or receipt upload is established in the main project.','Google sign-in has an Android implementation; the generic implementation reports Android availability.','Image uploads cover profiles and catalog imagery, not camera roll backup or client album delivery.','The app needs connectivity for authoritative booking operations. Cached browsing does not provide offline booking synchronization.','The draft is held in memory; full restart recovery is not established.','Status, history, audit and notification writes are not uniformly one atomic operation.','Public image buckets should not be presented as private photo storage.','Push submission is not delivery confirmation. Webhook configuration, native credentials and real devices require verification.'])
h('Future improvements to present as proposals')
p('Complete the Clients directory, add payments only after requirements are defined, centralize more business mutations in authorized transactional RPCs, unify all booking date calculations with studio time, consider secure credential storage, add push receipt and invalid token handling, and test on representative Android and iOS devices. Measure performance before adding more cache complexity.')

page('11 Defense questions and suggested answers')
qa=[
('What does PhotoSync solve','It centralizes photography service browsing, booking requests, studio availability and booking communication for clients and photographers.'),
('What is your backend','We use Supabase as the backend platform: Auth, PostgreSQL, row level security, storage, database functions and a Deno Edge Function for push notifications.'),
('Why use a relational database','The data has clear relationships among clients, services, packages, bookings and status history. Foreign keys, constraints and transactions help maintain those relationships.'),
('Why Expo and React Native','They provide a shared component based mobile application structure with navigation, images, notifications and build configuration in the same project. Cross-platform support still requires platform testing.'),
('How do you separate clients and admins','Profiles store a trusted role. Routing chooses the corresponding workspace, while database policies enforce ownership and admin permissions independently of the UI.'),
('Can a user register as an admin','Public signup creates a client profile. Admin promotion is performed through a trusted database operation.'),
('How do you prevent double booking','We recheck availability before submission and confirmation, condition writes on the current status, and use a database exclusion constraint to reject overlapping confirmed appointments.'),
('Are all pending requests guaranteed reservations','No. Submission creates a pending request. Final confirmation performs another availability check, and conflicting confirmed appointments are rejected by the database.'),
('What is preparation buffer','It is time around a session used by availability checks so studio preparation does not collide with another session or blocked interval.'),
('How do you handle changed prices','The app fetches the package again at confirmation. If the price changed, it asks the client to review and confirm the updated price.'),
('What happens to unanswered requests','An expiration RPC marks due pending requests expired using the server clock in Philippine studio time and records the transition atomically.'),
('Can clients cancel confirmed bookings','The current cancel action permits pending requests. A confirmed booking can enter the rescheduling request workflow and return to pending review.'),
('How are reminders created','A documented hourly database job selects tomorrow confirmed sessions in Manila time. Duplicate prevention protects against overlapping job and app reminder generation.'),
('How are push notifications sent','A notification record triggers the configured dispatcher. The Edge Function finds active recipient sessions and submits messages through Expo Push Service.'),
('How do you know a notification arrived','The current sender response shows submission results. It is not a guarantee of display on a device; real device delivery and receipt handling require separate verification.'),
('Why do you use caches','They reduce repeated reads and preserve screen content. Critical writes still fetch authoritative package, hours and availability data rather than trusting browsing freshness.'),
('What happens when users switch accounts','Account scoped caches clear and stale response guards prevent old requests from filling the new account state. Push registrations are also tied to active sessions.'),
('Is the app fully offline','No. Some previously read data can remain visible, but bookings and authoritative checks require the backend. It has no general offline write synchronization engine.'),
('Are uploaded photos private','The current avatar and catalog buckets have public reads and restricted writes. We should not describe them as private galleries.'),
('Does it include payments and a live Clients directory','Those are not completed features in the main application. The Clients screen is a placeholder; separate prototype capabilities should not be confused with the main backend.'),
('How did you validate the implementation','The current regression run passed 473 tests, and the TypeScript check passed. Database fixture and deployment evidence is recorded separately in the repository. Device tests remain necessary.'),
('What would you improve next','Complete unfinished operational screens, strengthen transactional booking mutations, unify timezone calculations, improve push delivery observability and test real device behavior.')]
for i,(q,a) in enumerate(qa,1):h(f'{i} {q}');p(a)

page('12 Demonstration and rehearsal plan')
h('Opening explanation')
p('PhotoSync is a mobile booking system for a photography studio. Clients can explore services, select a package and request a schedule. Photographers review requests and manage availability. The application uses Expo and React Native for the interface and Supabase for authentication, shared data and backend access controls.')
h('Suggested demonstration sequence')
steps=[('Prepare','Install the correct development build, confirm network access, prepare a client and an admin account, and verify published service and package records. Use demo contact information.'),('Show client login','Sign in as a client and explain the role based landing page. Demonstrate signup or recovery only if its external configuration is ready.'),('Browse catalog','Show a service and package. Explain duration, price, notice and preparation rules.'),('Create a request','Pick a valid future date, choose a slot, enter contact information and review the request. Submit once and explain the pending state.'),('Review as photographer','Sign out and enter the admin workspace. Open the submitted request, inspect its details and confirm or reject it.'),('Show client result','Return to the client and refresh the booking view or inbox. Demonstrate the resulting status. Treat push as an additional demo only after verifying the device setup.'),('Show availability','Open a future calendar day and explain available versus unavailable windows. Show that a past date is read only.'),('Show catalog management','Explain publication versus archive and how older booking references are preserved. Avoid changing business records merely for demonstration.'),('Explain backend evidence','Show the entity relationships, role policies and confirmed overlap constraint without exposing credentials or private client records.'),('Close with scope','Name the completed booking workflow and the specific remaining features. State testing evidence accurately.')]
for i,(t,b) in enumerate(steps,1):h(f'{i} {t}');p(b)
h('Before leaving for the defense')
bullets(['Charge the device and prepare a backup network connection.','Test client and admin credentials and the exact build you will demonstrate.','Choose a future date satisfying the package notice requirement.','Verify notification permissions and Google credentials if demonstrating those features.','Keep a screenshot or screen recording of the successful flow as a fallback.','Rehearse the difference between implemented behavior, historical verification and future proposals.'])

page('Appendix A Source map')
p('Repository locations below identify the implementation evidence for studying and answering follow-up questions. Paths are relative to C:/Users/kirlg/PhotoSync unless explicitly named otherwise. No secret values are included.')
sources=[('Stack and native configuration','package.json; app.json; eas.json'),('Auth and role routing','src/lib/supabase.ts; src/services/auth.ts; src/hooks/use-auth-routing.ts; src/app/_layout.tsx'),('Google login','src/services/google-auth.android.ts; src/services/google-auth.ts; docs/google-sign-in.md'),('Database structure and setup','docs/supabase-schema.sql; docs/supabase-setup.md; docs/supabase-rls-hardening.sql'),('Submission and rules','src/services/bookings.ts; src/services/booking-availability.ts; src/services/booking-draft.ts'),('Admin and client actions','src/services/admin-bookings.ts; src/services/client-bookings.ts'),('Confirmed conflict protection','docs/supabase-confirmed-booking-conflict-protection.sql; docs/supabase-booking-conflicts.sql'),('Calendar and timezone','src/services/calendar.ts; src/services/calendar-date-guards.ts; docs/calendar-date-guards.md'),('Catalog and archive','src/services/service-catalog.ts; docs/catalog-deletion.md; docs/supabase-catalog-archives.sql'),('Profiles and image storage','src/services/profile.ts; docs/supabase-profile-images.sql; docs/supabase-service-images.sql'),('Notifications and push','src/services/notifications.ts; src/services/push-notifications.ts; supabase/functions/send-push-notification/index.ts'),('Reminders','docs/supabase-scheduled-booking-reminders.sql; docs/supabase-reminder-schedule.sql'),('Pagination and expiration','src/services/booking-pages.ts; docs/booking-pagination.md; docs/booking-expiration.md'),('Shared performance work','src/services/session-read-cache.ts; docs/shared-service-optimization.md'),('Placeholder evidence','src/app/photographer/clients.tsx'),('Diagrams available in repository','docs/PhotoSync_ERD.drawio; docs/PhotoSync_Booking_Flow.drawio'),('Separate local prototype','C:/Users/kirlg/PhotoSyncMobile/README.md; its src/data modules')]
for name,loc in sources:h(name);p(loc)

page('Appendix B SQL routine inventory')
p('This inventory lists function definitions found in the repository SQL files. A definition in source does not by itself establish that the matching routine is installed in the deployed database. Repeated definitions are grouped under the same routine name.')
funcs={}
for file in sorted((ROOT/'docs').glob('supabase*.sql')):
 for name in re.findall(r'create\s+(?:or\s+replace\s+)?function\s+public\.([a-zA-Z0-9_]+)',file.read_text(encoding='utf-8-sig'),re.I):
  funcs.setdefault(name,[]).append(file.name)
table(['Routine','SQL source files'],[(k,'; '.join(v)) for k,v in sorted(funcs.items())],[2.7,4.2])
d.core_properties.title='PhotoSync Project Defense Guide'
d.core_properties.subject='Backend and mobile implementation with defense preparation'
d.core_properties.author='PhotoSync Project Team'
target=OUT/'PhotoSync_Project_Defense_Guide.docx'
d.save(target)
print(target)
print('Words',sum(len(x.text.split()) for x in d.paragraphs))

